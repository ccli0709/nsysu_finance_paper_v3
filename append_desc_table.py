# -*- coding: utf-8 -*-
"""在 中山格式 論文檔最後，加上與樣版「表4-3 敘述統計」相同格式之表格（本研究資料）。"""
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = "水與廢棄物管理與成本黏性_論文_中山格式.docx"
MODEL_CSV = "04_model_data.csv"
CJK, EN = "標楷體", "Times New Roman"

STATS = ["Mean", "SD", "Min", "P25", "Median", "P75", "Max"]
DESC = [("Y_dLNSGA", "ΔLNSGA"), ("dLNREV", "ΔLNREV"), ("D", "D"),
        ("Water_Intensity", "用水密集度"), ("Waste_Intensity", "廢棄物密集度"),
        ("Water_Rate", "水回收率%"), ("Waste_Disc", "廢棄物揭露度"),
        ("Waste_Fine", "廢棄物裁罰金額"),
        ("Size", "Size"), ("AI", "AI"), ("EI", "EI"), ("ROA", "ROA"),
        ("Lev", "Lev"), ("Decrease", "Decrease")]


def font(run, size=12, bold=False):
    run.font.name = EN
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)


def set_border(cell, edge, sz=8):
    tcPr = cell._tc.get_or_add_tcPr()
    bd = tcPr.find(qn("w:tcBorders"))
    if bd is None:
        bd = OxmlElement("w:tcBorders"); tcPr.append(bd)
    e = bd.find(qn(f"w:{edge}"))
    if e is None:
        e = OxmlElement(f"w:{edge}"); bd.append(e)
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")


def fill(cell, text, bold=False, align="right"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    font(p.add_run(text), 12, bold)


df = pd.read_csv(MODEL_CSV, encoding="utf-8-sig", low_memory=False)
doc = Document(DOC)
doc.add_page_break()
cap = doc.add_paragraph()
font(cap.add_run("表4-3 敘述統計"), 12, True)

tbl = doc.add_table(rows=1 + len(DESC), cols=1 + len(STATS))
tbl.style = "Normal Table"
fill(tbl.rows[0].cells[0], "", align="left")
for j, s in enumerate(STATS):
    fill(tbl.rows[0].cells[j + 1], s, align="right")
for i, (col, lab) in enumerate(DESC):
    s = pd.to_numeric(df[col], errors="coerce").dropna()
    vals = [s.mean(), s.std(), s.min(), s.quantile(.25), s.median(), s.quantile(.75), s.max()]
    fill(tbl.rows[i + 1].cells[0], lab, align="left")
    for j, v in enumerate(vals):
        fill(tbl.rows[i + 1].cells[j + 1], f"{v:.2f}", align="right")
for c in tbl.rows[0].cells:
    set_border(c, "top", 8); set_border(c, "bottom", 8)
for c in tbl.rows[-1].cells:
    set_border(c, "bottom", 8)

doc.save(DOC)
print("已於文末加入表4-3 敘述統計（樣版格式）：", DOC)
